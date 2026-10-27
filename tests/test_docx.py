"""The DOCX booklet: presentation text, clue wording and embedded diagrams."""

import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document

from csp.cli import main
from csp.docx_export import Entry, write_booklet
from csp.puzzlefile import load_presentation
from csp.story import clue_paragraphs, clue_texts, object_list
from tests.examples import all_examples, example

PRRRDOKUS = ["prrrdoku1", "prrrdoku2", "prrrdoku3"]


def write_docx(text, out, base, solution=None):
    write_booklet([Entry(text, base / "x.txt", solution)], out)


def _text(path: Path) -> str:
    document = Document(str(path))
    cells = [p for t in document.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    return "\n".join(p.text for p in [*document.paragraphs, *cells])


@pytest.mark.parametrize("name", PRRRDOKUS)
def test_prrrdoku_booklet_has_text_and_diagram(name, tmp_path):
    source = example(f"{name}.txt")
    out = tmp_path / f"{name}.docx"
    assert main([str(source), "--docx", str(out)]) == 0
    text = _text(out)
    assert "Prrrdoku" in text and "Aanwijzingen" in text and "Vladimir" in text
    assert any(n.startswith("word/media/") for n in zipfile.ZipFile(out).namelist())
    assert "Oplossing" in text
    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("word/media/")]
    assert len(media) == 2 + len(object_list(source.read_text()))  # puzzle, solution, objects
    assert text.index("Aanwijzingen") < text.index("Oplossing")


def test_clue_text_override_wins():
    source = (example("prrrdoku1.txt")).read_text()
    assert load_presentation(source).clue_text[5].startswith("Otto zit")


def test_puzzle_without_presentation_still_gets_a_booklet(tmp_path):
    source = example("murdoku_house.txt")
    plain = source.read_text().split("# Booklet text")[0]
    assert not load_presentation(plain).title
    out = tmp_path / "plain.docx"
    write_docx(plain, out, source.parent)
    text = _text(out)
    assert "Clues" in text
    assert all(paragraph in text for paragraph in clue_paragraphs(plain))


def test_missing_figure_image_falls_back_to_text(tmp_path):
    base = (example("prrrdoku1.txt")).read_text()
    source = base.replace(
        "Figures:\nVladimir: de poes die iemand op schoot zoekt",
        "Figures:\nVladimir: de poes ; image=missing.png",
    )
    out = tmp_path / "x.docx"
    write_docx(source, out, tmp_path)
    assert "Vladimir: de poes" in _text(out)


def test_pages_pdf_export_converts_paths_to_posix_files(monkeypatch, tmp_path):
    source = tmp_path / "booklet.docx"
    target = tmp_path / "booklet.pdf"
    source.write_bytes(b"docx")

    called = {}

    def fake_run(command):
        called["command"] = command
        target.write_bytes(b"pdf")
        return type("Result", (), {"returncode": 0, "stdout": "", "stderr": ""})()

    monkeypatch.setattr("csp.pdf_convert.converters", lambda: ["Pages"])
    monkeypatch.setattr("csp.pdf_convert._run", fake_run)

    from csp.pdf_convert import docx_to_pdf

    assert docx_to_pdf(source, target) == "Pages"
    script = called["command"][2]
    assert "set sourceFile to POSIX file source" in script
    assert "set targetFile to POSIX file target" in script
    assert "export theDocument to targetFile as PDF" in script
    assert target.exists()


def test_figure_image_is_embedded(tmp_path):
    from PIL import Image

    Image.new("RGB", (20, 20), "red").save(tmp_path / "v.png")
    source = (
        (example("prrrdoku1.txt"))
        .read_text()
        .replace("Vladimir: de poes die iemand op schoot zoekt", "Vladimir: poes ; image=v.png")
    )
    out = tmp_path / "x.docx"
    write_docx(source, out, tmp_path)
    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("word/media/")]
    assert len(media) == 2 + len(object_list(source))


def test_bad_presentation_is_rejected_and_linted():
    bad = (example("prrrdoku1.txt")).read_text() + "\nFigures:\nNobody: x\n"
    with pytest.raises(ValueError):
        load_presentation(bad)


def test_human_readable_clues():
    got = clue_texts((example("food-murdoku-3.txt")).read_text())
    assert "Okonomiyaki is in column 1." in got
    assert any("somewhere above" in clue for clue in got)


def test_glob_makes_one_multi_puzzle_booklet(tmp_path):
    for name in PRRRDOKUS[:2]:
        (tmp_path / f"{name}.txt").write_text((example(f"{name}.txt")).read_text())
    out = tmp_path / "all.docx"
    assert main([str(tmp_path / "prrrdoku*.txt"), "--docx", str(out)]) == 0
    text = _text(out)
    assert "Prrrdoku 1" in text and "Prrrdoku 2" in text
    assert text.count("Aanwijzingen") == 2 and "Oplossing" in text
    assert text.rindex("Prrrdoku 1") > text.index("Oplossing")
    assert not (tmp_path / "all-solution.docx").exists()


def test_inline_wording_does_not_change_the_puzzle():
    from csp.puzzlefile import load_board

    plain = (
        (example("prrrdoku1.txt"))
        .read_text()
        .replace(" | Otto zit precies één rij boven Tjitske.", "")
    )
    assert load_board(plain) == load_board((example("prrrdoku1.txt")).read_text())


def _media(path):
    return [n for n in zipfile.ZipFile(path).namelist() if n.startswith("word/media/")]


def test_figures_folder_gives_thumbnails_by_name(tmp_path):
    from PIL import Image

    text = (example("tamagotchi-murdoku-6x6.txt")).read_text()
    (tmp_path / "t.txt").write_text(text)
    folder = tmp_path / "t-figures"
    folder.mkdir()
    Image.new("RGB", (20, 20), "green").save(folder / "Mametchi.png")
    out = tmp_path / "x.docx"
    write_booklet([Entry(text, tmp_path / "t.txt")], out)
    assert len(_media(out)) > 7  # board, six figures, then the objects


def test_icon_figures_use_builtin_icons(tmp_path):
    # Every icon= in the example is vendored, so each figure gets a thumbnail.
    text = (example("food-murdoku-1.txt")).read_text()
    out = tmp_path / "x.docx"
    write_booklet([Entry(text, example("food-murdoku-1.txt"))], out)
    figures = sum(1 for line in text.splitlines() if "icon=" in line)
    assert figures >= 4
    assert len(_media(out)) >= 1 + figures


def test_unknown_icon_is_linted(tmp_path, capsys):
    text = (example("food-murdoku-1.txt")).read_text().replace("U+1F35C", "nonsense")
    (tmp_path / "f.txt").write_text(text)
    main([str(tmp_path / "f.txt"), "--lint"])
    assert "unknown icon nonsense" in capsys.readouterr().out


def test_icon_that_is_not_vendored_is_linted(tmp_path, capsys):
    # A valid code point whose picture was never vendored would print as text.
    text = (example("food-murdoku-1.txt")).read_text().replace("U+1F35C", "U+1F680")
    (tmp_path / "f.txt").write_text(text)
    main([str(tmp_path / "f.txt"), "--lint"])
    assert "unknown icon U+1F680" in capsys.readouterr().out


def test_booklet_footer_carries_copyright_notice(tmp_path):
    from docx import Document

    from csp.docx_export import COPYRIGHT_NOTICE

    out = tmp_path / "b.docx"
    assert main([str(example("prrrdoku1.txt")), "--docx", str(out)]) == 0
    assert COPYRIGHT_NOTICE in Document(str(out)).sections[0].footer.paragraphs[0].text


@pytest.mark.parametrize("name", ["sudoku_easy", "queens_4x4", "tents_4x4", "calcudoku_4x4_easy"])
def test_other_families_get_rules_puzzle_and_solution(name, tmp_path):
    out = tmp_path / "b.docx"
    assert main([str(example(f"{name}.txt")), "--docx", str(out)]) == 0
    document = Document(str(out))
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert len(headings) == 2 and headings[1].startswith("Solution")
    assert any(p.text == "Rules" for p in document.paragraphs)
    assert any(p.text == "How to solve it" for p in document.paragraphs)
    assert len(document.inline_shapes) == 2


def _read(name):
    return (example(f"{name}.txt")).read_text()


def test_clues_are_grouped_per_person():
    paragraphs = clue_paragraphs(_read("zoo-murdoku-1"))
    makaak = next(p for p in paragraphs if p.startswith("De makaak"))
    assert makaak == "De makaak is in de bergen en zit niet op of naast het gras."
    assert sum(p.startswith("De giraffe") for p in paragraphs) == 1


def test_exactly_one_same_region_reads_as_either_or():
    paragraphs = clue_paragraphs(_read("murdoku_sparse_mystery"))
    assert any("same region as either Mara or Ria" in p for p in paragraphs)


def test_two_alternatives_read_as_either_or_not_both():
    clues = clue_texts(_read("murdoku_house"))
    assert any(c.startswith("Either Ada is in the study, or Ben is in the hall") for c in clues)


def test_attribute_target_can_be_reworded():
    text = _read("tamagotchi-murdoku-12x12").replace(" Ichigotchi Hashizotchi", " Ichigotchi", 1)
    assert (
        "Akumatchi shares a region with exactly one other person, and that is a friend."
        in clue_paragraphs(text)
    )


def test_attribute_everyone_has_is_not_named():
    paragraphs = clue_paragraphs(_read("tamagotchi-murdoku-12x12"))
    assert "Akumatchi shares a region with exactly one other person." in paragraphs


def test_solution_without_clue_numbers():
    from csp.story import explain

    assert (
        "(" not in explain(_read("zoo-murdoku-1"), numbered=False).replace("**", "").split(":")[0]
    )
    assert "which clue" not in explain(_read("murdoku_composition"), numbered=False)


def test_zoo_roles_are_not_called_roles():
    from csp.story import explain

    text = explain(_read("zoo-murdoku-1"), numbered=False)
    assert "de rol" not in text
    # The figure list already says which animals are apes; the booklet does not repeat it.
    assert "is geen aap" not in text
    assert "is geen aap" in explain(_read("zoo-murdoku-1"))


def test_no_example_has_an_implied_clue():
    from csp.murdoku import implied_clues
    from csp.puzzlefile import detect, load_board

    for path in all_examples():
        text = path.read_text()
        if detect(text) == "murdoku":
            assert not implied_clues(load_board(text)[1]), path.name


def _lint_output(text):
    from csp.cli import _murdoku_lint
    from csp.core import Model  # noqa: F401
    from csp.murdoku import compile_puzzle
    from csp.puzzlefile import load_board

    board, clues = load_board(text)
    model = compile_puzzle(board, clues)
    return "\n".join(_murdoku_lint(board, clues, model, text))


def test_zoo_puzzles_have_no_one_person_role():
    for i in (1, 2, 3):
        assert "can only be held" not in _lint_output(_read(f"zoo-murdoku-{i}"))


def test_single_holder_role_is_linted():
    text = _read("zoo-murdoku-1").replace("RoleEligibility:\n", "RoleEligibility:\nsolo: Leeuw\n")
    text = text.replace("Roles:\n", "Roles:\nsolo: 1\n")
    assert "role solo can only be held by Leeuw" in _lint_output(text)


def test_furthest_clue_names_nobody_to_compare_with():
    text = _read("prrrdoku3")
    assert "van alle figuren het verst van Mao" in " ".join(clue_paragraphs(text))
    assert "dan Tim" not in " ".join(clue_paragraphs(text))


def test_board_notes_explain_marks_and_attributes_and_roles_are_listed_apart():
    from csp.story import attribute_list, board_notes, role_list

    notes = board_notes(_read("murdoku_composition"))
    assert any("whoever stands on one is alone in their area" in n for n in notes)
    assert not any("trained" in n for n in notes)
    assert ("trained", "Cal and Dee") in attribute_list(_read("murdoku_composition"))
    assert not any("navigator" in n for n in notes)
    assert ("navigator", "held by exactly one person, chosen from Cal and Dee") in role_list(
        _read("murdoku_composition")
    )


def test_booklet_lists_roles_under_the_board(tmp_path):
    out = tmp_path / "r.docx"
    assert main([str(example("murdoku_composition.txt")), "--docx", str(out)]) == 0
    texts = [p.text for p in Document(str(out)).paragraphs]
    assert any(t.startswith("The attributes") and "trained: Cal and Dee" in t for t in texts)
    assert any(t.startswith("The roles") and "navigator: held by" in t for t in texts)


def test_booklet_groups_objects_by_kind_on_one_line(tmp_path):
    out = tmp_path / "o.docx"
    assert main([str(example("murdoku_orient_express.txt")), "--docx", str(out)]) == 0
    line = next(p.text for p in Document(str(out)).paragraphs if p.text.startswith("De voorwerpen"))
    assert "Begaanbaar:" in line and "Blokkerend:" in line
    assert line.index("Treinbank") < line.index("Blokkerend:") < line.index("Klok")
    assert "hier kan niemand staan" not in line


def test_booklet_numbers_the_clues_and_lists_the_figures_beside_the_board(tmp_path):
    out = tmp_path / "b.docx"
    assert main([str(example("prrrdoku1.txt")), "--docx", str(out)]) == 0
    document = Document(str(out))
    assert next(p.text for p in document.paragraphs if p.text[:2] == "1.").startswith("1.\t")
    assert len(document.tables) == 1
    names = [p.text for p in document.tables[0].rows[0].cells[1].paragraphs]
    assert any(n.startswith("Vladimir") for n in names)


def test_board_notes_name_the_areas_of_a_tag():
    from csp.story import board_notes

    notes = board_notes(_read("murdoku_sparse_mystery"))
    assert "The noble areas are Maras house and Blue hall." in notes


def test_booklet_is_numbered_by_category_and_has_a_contents_page(tmp_path):
    out = tmp_path / "t.docx"
    names = [str(example("food-murdoku-1.txt")), str(example("zoo-murdoku-1.txt"))]
    assert main([*names, "--docx", str(out)]) == 0
    document = Document(str(out))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "Inhoud"  # the first puzzle, Dutch, sets the language
    lines = [t.split("\t")[0] for t in texts if "\t" in t]
    assert [line.split(":")[0] for line in lines[:3]] == [
        "1. Dierentuin-murdoku 5x5",
        "2. Food Murdoku 5x5",
        "Oplossingen",
    ]
    assert "Murdoku \u00b7 Japanese food" in texts
    assert any(text.startswith("Solution 2") for text in texts)
    xml = document.element.xml
    assert 'w:anchor="puzzle0"' in xml and 'w:name="puzzle1"' in xml


def test_single_puzzle_booklet_has_no_contents(tmp_path):
    out = tmp_path / "s.docx"
    assert main([str(example("prrrdoku1.txt")), "--docx", str(out)]) == 0
    assert "Contents" not in [p.text for p in Document(str(out)).paragraphs]


def test_booklet_solutions_have_no_empty_reasons_or_clue_numbers():
    import re

    from csp.story import explain

    for name in ("menu-murdoku-3", "zoo-murdoku-1", "food-murdoku-2", "prrrdoku1"):
        text = explain(_read(name), numbered=False)
        assert ", ." not in text and ": ." not in text, name
        assert not re.search(r"\b(clue|aanwijzing) \d", text), name
        assert not re.search(r"\(via [\d, en]+\)", text), name


def test_repeated_premises_are_merged():
    import itertools
    import re

    from csp.story import explain

    for name in ("menu-murdoku-3", "zoo-murdoku-1", "zoo-murdoku-3", "prrrdoku2"):
        text = explain(_read(name), numbered=False)
        premises = [
            re.split(", (?:dus|so) ", s)[0] if re.search(", (?:dus|so) ", s) else None
            for s in re.split(r"(?<=\.) ", text)
        ]
        repeats = [a for a, b in itertools.pairwise(premises) if a is not None and a == b]
        assert not repeats, (name, repeats)


def test_pdf_writes_every_example_murdoku_without_puzzle_files(tmp_path, monkeypatch, capsys):
    import shutil

    from csp import cli, pdf_convert

    seen = {}

    def fake_convert(source, target):
        seen["text"] = "\n".join(p.text for p in Document(str(source)).paragraphs)
        shutil.copy(source, target)
        return "Fake Office"

    monkeypatch.setattr(pdf_convert, "docx_to_pdf", fake_convert)
    monkeypatch.chdir(tmp_path)
    files = cli._murdoku_examples()
    assert files and all("tests" not in file.parts for file in files)
    assert {file.parent.name for file in files} >= {"demo", "prrrdoku", "wie-doet-het-met-wie"}
    assert main(["--pdf"]) == 0
    assert (tmp_path / "booklet.pdf").exists()
    assert "Prrrdoku 1" in seen["text"] and "Sudoku" not in seen["text"]
    assert f"({len(files)} puzzles" in capsys.readouterr().out


def test_pdf_with_a_puzzle_file_writes_beside_it(tmp_path, monkeypatch):
    import shutil

    from csp import pdf_convert

    monkeypatch.setattr(
        pdf_convert, "docx_to_pdf", lambda source, target: shutil.copy(source, target) and "Fake"
    )
    source = tmp_path / "prrrdoku1.txt"
    shutil.copy(example("prrrdoku1.txt"), source)
    assert main([str(source), "--pdf"]) == 0
    assert (tmp_path / "prrrdoku1.pdf").exists()


def test_pdf_without_an_office_program_explains_what_to_install(tmp_path, monkeypatch):
    from csp import pdf_convert

    monkeypatch.setattr(pdf_convert, "converters", lambda: [])
    with pytest.raises(pdf_convert.ConversionError, match="install LibreOffice"):
        pdf_convert.docx_to_pdf(tmp_path / "x.docx", tmp_path / "x.pdf")


def test_docx_without_puzzle_files_takes_the_example_murdoku(tmp_path, monkeypatch, capsys):
    # The PDF test above builds the booklet of every example; two puzzles are
    # enough to check the default --docx name and count.
    from csp import cli

    files = [file.resolve() for file in cli._murdoku_examples()[:2]]
    monkeypatch.setattr(cli, "_murdoku_examples", lambda: files)
    monkeypatch.chdir(tmp_path)
    assert main(["--docx"]) == 0
    assert (tmp_path / "booklet.docx").exists()
    assert "(2 puzzles" in capsys.readouterr().out


def test_an_object_with_several_pieces_is_named_indefinitely_despite_its_words():
    clues = clue_texts(_read("murdoku_orient_express"))
    assert "Anastasia zit op een perronbankje." in clues
    assert "Vladimir zit niet op of naast een treinbank." in clues
    assert not any("het perronbankje" in clue for clue in clues)


def test_english_uses_an_before_a_vowel_for_one_of_several_objects():
    from csp.story import Places, Wording, load_board

    text = """Size: 3
Language: en
Regions:
a: room #EEEEEE
Grid:
a a a
a a a
a a a
People:
Ada
Objects:
apple: standable r1c1
apple: standable r3c3
bench: standable r1c3
bench: standable r3c1
Words:
bench: the bench
"""
    board, _ = load_board(text)
    places = Places(board, Wording.read(text))
    assert places.thing("apple") == "an apple"
    assert places.thing("bench") == "a bench"


def test_one_of_several_windows_is_named_indefinitely():
    clues = clue_texts(_read("murdoku_orient_express_3"))
    assert "Bogdan zit naast een raam." in clues
    assert not any("het raam" in clue for clue in clues)
