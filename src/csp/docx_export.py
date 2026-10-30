"""An editable Word booklet page for a Murdoku puzzle file (needs python-docx).

The booklet reads the optional presentation sections of the puzzle file
(`Title:`, `Intro:`, `Question:`, `ClueText:`, `Figures:`, `Reveal:`), draws the
board with `picture` and words the clues with `story`. Nothing here is needed to
solve; a file without presentation sections still gives a usable booklet.
"""

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from . import icons, picture
from .puzzlefile import Figure, Presentation, language_of, load_board, load_presentation
from .story import (
    attribute_list,
    board_notes,
    clue_paragraphs,
    explain,
    family_rules,
    object_list,
    phrase,
    role_list,
)

BOLD = re.compile(r"\*\*(.+?)\*\*")
BOARD_WIDTH_CM = 11.6
MIN_BOARD_WIDTH_CM = 5.0
TEXT_WIDTH_CM = 17.4
BOARD_SHARE = 0.7
PAGE_BODY_CM = 26.4
FIXED_CM = 7.8
RULE_LINE_CM = 0.48
CLUE_LINE_CM = 0.5
SMALL_CLUE_LINE_CM = 0.4
MANY_CLUES = 18
CLUE_LINE_CHARS = 105
MANY_SOLUTION_LINES = 30
SOLUTION_MIN_BOARD_CM = 5.0
FIGURE_LINE_CM = 1.2
FIGURE_IMAGE_CM = 0.7
MIN_FIGURE_IMAGE_CM = 0.4
FIGURE_HEADING_CM = 0.7
FIGURE_GAP_CM = 0.27
OBJECT_IMAGE_CM = 0.5
GENERIC_BOARD_WIDTH_CM = 14.0


def _runs(paragraph: Any, text: str) -> None:
    """Add `text` to a paragraph, turning **bold** markers into bold runs."""
    position = 0
    for match in BOLD.finditer(text):
        paragraph.add_run(text[position : match.start()])
        paragraph.add_run(match[1]).bold = True
        position = match.end()
    paragraph.add_run(text[position:])


def _picture_stream(
    text: str,
    solution: dict[str, Any] | None,
    role_statuses: dict[tuple[str, str], bool | None] | None,
    backdrop: bool = True,
) -> tuple[io.BytesIO, float]:
    """The board picture and its height/width ratio; `backdrop=False` leaves the
    backdrop out, for pages too full to give it room."""
    board, _ = load_board(text)
    if not backdrop:
        board.backdrop = None
    image = picture.murdoku_picture(
        board, solution, show_roles=False, language=language_of(text), area_tags=False
    )
    return io.BytesIO(image.to_png()), image.height / image.width


def _puzzle_board_width(aspect: float, clue_lines: int, clues: int, rules_lines: int) -> float:
    """Board width in cm, shrunk so the rules, clues and question still fit the page."""
    line = CLUE_LINE_CM if clue_lines <= MANY_CLUES else SMALL_CLUE_LINE_CM
    free = PAGE_BODY_CM - FIXED_CM - RULE_LINE_CM * rules_lines - line * clue_lines - 0.12 * clues
    return max(MIN_BOARD_WIDTH_CM, min(BOARD_WIDTH_CM, free / aspect))


def _figure_image(person: str, figure: Figure | None, source: Path) -> Path | io.BytesIO | None:
    """The thumbnail for a person: the `image=` file, a file in `<stem>-figures/` or
    `figures/` beside the puzzle named after the person, or the built-in `icon=`."""
    base = source.parent
    if figure is not None and figure.image and (base / figure.image).is_file():
        return base / figure.image
    for folder in (base / f"{source.stem}-figures", base / "figures"):
        for suffix in (".png", ".jpg", ".jpeg"):
            if (folder / f"{person}{suffix}").is_file():
                return folder / f"{person}{suffix}"
    if figure is not None and figure.icon:
        icon = icons.object_icon(figure.icon)
        thumbnail = None if icon is None else icons.png(icon, 128)
        return None if thumbnail is None else io.BytesIO(thumbnail)
    return None


def _figure_image_cm(board_height: float, people: int) -> float:
    """Thumbnail height that keeps the list of `people` no taller than the board beside it."""
    fit = (board_height - FIGURE_HEADING_CM) / max(1, people) - FIGURE_GAP_CM
    return max(MIN_FIGURE_IMAGE_CM, min(FIGURE_IMAGE_CM, fit))


def _figure_list(
    cell: Any,
    text: str,
    presentation: Presentation,
    source: Path,
    image_cm: float = FIGURE_IMAGE_CM,
) -> None:
    """One paragraph per person, with the thumbnail and the description if there are any."""
    from docx.shared import Cm, Pt

    board, _ = load_board(text)
    figures = presentation.figures
    heading = cell.paragraphs[0]
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(phrase(text, "doc.figures"))
    _section_run(run)
    for person in board.people:
        figure = figures.get(person)
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        image = _figure_image(person, figure, source)
        if image is not None:
            source_image = str(image) if isinstance(image, Path) else image
            paragraph.add_run().add_picture(source_image, height=Cm(image_cm))
            paragraph.add_run(" ")
        paragraph.add_run(person).bold = True
        if figure is not None and figure.description:
            paragraph.add_run(f": {figure.description}")


def _list_lines(text: str) -> int:
    """Rule-sized lines that `_board_lists` takes; object lines count extra for their icons."""
    items = [f"{name}: {description}" for name, description in attribute_list(text)]
    items += [f"{name}: {description}" for name, description in role_list(text)]
    lines = sum(-(-(len(item) + 20) // CLUE_LINE_CHARS) for item in items)
    objects = object_list(text)
    if objects:
        chars = 50 + sum(len(label) + 8 for _, label, _ in objects)
        lines += 2 * -(-chars // CLUE_LINE_CHARS)
    return lines


def _board_lists(document: Any, text: str) -> None:
    """Attributes, roles and objects under the board, one compact line per item or group."""
    from docx.shared import Cm, Pt

    def line(title: str) -> Any:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        _section_run(paragraph.add_run(title + "  "))
        return paragraph

    for key, items in (("doc.attributes", attribute_list(text)), ("doc.roles", role_list(text))):
        for index, (name, description) in enumerate(items):
            paragraph = line(phrase(text, key)) if index == 0 else document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(name).bold = True
            paragraph.add_run(f": {description}")
    objects = object_list(text)
    # Objects are grouped by kind and flow side by side, so a long list stays short.
    paragraph = line(phrase(text, "doc.objects")) if objects else None
    for standable in (True, False):
        group = [(kind, label) for kind, label, can in objects if can == standable]
        if not group or paragraph is None:
            continue
        if not standable and any(can for _, _, can in objects):
            paragraph.add_run("   ")
        paragraph.add_run(
            phrase(text, "doc.standable_group" if standable else "doc.blocking_group") + ": "
        ).italic = True
        for index, (kind, label) in enumerate(group):
            if index:
                paragraph.add_run(", ")
            icon = icons.object_icon(kind)
            thumbnail = None if icon is None else icons.png(icon, 128)
            if thumbnail is not None:
                paragraph.add_run().add_picture(io.BytesIO(thumbnail), height=Cm(OBJECT_IMAGE_CM))
                paragraph.add_run("\u00a0")
            # Non-breaking spaces keep a multi-word name on one line with its icon.
            paragraph.add_run(label.replace(" ", "\u00a0"))


@dataclass(frozen=True)
class Entry:
    """One puzzle of a booklet; `solution` makes its page a worked solution."""

    text: str
    path: Path
    solution: dict[str, Any] | None = None
    role_statuses: dict[tuple[str, str], bool | None] | None = None
    kind: str = "murdoku"
    puzzle_image: Any = None
    solution_image: Any = None
    steps: tuple[str, ...] = ()
    number: int = 0
    section: str = ""

    def numbered(self, title: str) -> str:
        """`title` with the puzzle's number in front, if it has one."""
        return f"{self.number}. {title}" if self.number else title


def _generic_title(entry: Entry) -> str:
    return load_presentation(entry.text).title or entry.path.stem.replace("_", " ").capitalize()


def _add_generic(
    document: Any,
    entry: Entry,
    title: str,
    new_page: bool,
    solved: bool,
    bookmark: str | None = None,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    _title(document, title, new_page, bookmark, entry.section)
    if not solved:
        document.add_heading(phrase(entry.text, "doc.rules_title"), level=2)
        for rule in family_rules(entry.text, entry.kind):
            document.add_paragraph(rule, style="List Bullet")
    image = entry.solution_image if solved else entry.puzzle_image
    stream = io.BytesIO(image.to_png(zoom=2))
    document.add_picture(stream, width=Cm(GENERIC_BOARD_WIDTH_CM))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if solved and entry.steps:
        document.add_heading("How to solve it", level=2)
        for line in entry.steps:
            paragraph = document.add_paragraph(line, style="List Bullet")
            for run in paragraph.runs:
                run.font.size = Pt(8)


def _title(
    document: Any, title: str, new_page: bool, bookmark: str | None = None, section: str = ""
) -> None:
    """The big spaced title and, after a colon, an italic subtitle.

    A `bookmark` name makes the title a target for the table of contents; a
    `section` goes above it in small grey capitals.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    main, _, sub = title.partition(": ")
    if section:
        kicker = document.add_paragraph()
        kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
        kicker.paragraph_format.page_break_before = new_page
        kicker.paragraph_format.space_after = Pt(0)
        kicker.paragraph_format.keep_with_next = True
        run = kicker.add_run(section)
        run.font.all_caps = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        new_page = False
    heading = document.add_heading(main, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.page_break_before = new_page
    if bookmark:
        _bookmark(heading, bookmark)
    if sub:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(sub)
        run.italic = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _boxed(document: Any, text: str) -> None:
    """A bold, centred line in a box: the goal of the puzzle."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.left_indent = paragraph.paragraph_format.right_indent = Pt(8)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "8"), ("space", "8"), ("color", "888888")):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    _add_borders(paragraph._p.get_or_add_pPr(), borders)
    _runs(paragraph, text)
    for run in paragraph.runs:
        run.bold = True


def _numbered(document: Any, number: int, text: str, small: bool) -> None:
    """A clue with its number and a hanging indent."""
    from docx.shared import Cm, Pt

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Cm(0.7)
    paragraph.paragraph_format.first_line_indent = Cm(-0.7)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
    paragraph.add_run(f"{number}.\t")
    _runs(paragraph, text)
    if small:
        for run in paragraph.runs:
            run.font.size = Pt(9)


def _board_and_figures(
    document: Any, entry: Entry, presentation: Presentation, width: float, backdrop: bool = True
) -> None:
    """The board at 70% of the text width, the list of people in the other 30%."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    stream, aspect = _picture_stream(entry.text, None, None, backdrop)
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Cm(TEXT_WIDTH_CM * BOARD_SHARE)
    right.width = Cm(TEXT_WIDTH_CM * (1 - BOARD_SHARE))
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    left.paragraphs[0].add_run().add_picture(stream, width=Cm(width))
    people = len(load_board(entry.text)[0].people)
    image_cm = _figure_image_cm(width * aspect, people)
    _figure_list(right, entry.text, presentation, entry.path, image_cm)


def _puzzle_title(entry: Entry) -> str:
    if entry.kind != "murdoku":
        return entry.numbered(_generic_title(entry))
    first_comment = next((line for line in entry.text.splitlines() if line.startswith("# ")), "")
    title = load_presentation(entry.text).title
    return entry.numbered(title or first_comment.removeprefix("# ").strip() or entry.path.stem)


def _solution_title(entry: Entry) -> str:
    title = _generic_title(entry) if entry.kind != "murdoku" else _murdoku_title(entry)
    label = "Solution" if entry.kind != "murdoku" else phrase(entry.text, "doc.solution")
    if entry.number:
        label = f"{label} {entry.number}"
    return f"{label}: {title}"


def _murdoku_title(entry: Entry) -> str:
    return load_presentation(entry.text).title or entry.path.stem


def _add_puzzle(document: Any, entry: Entry, new_page: bool, bookmark: str | None = None) -> None:
    if entry.kind != "murdoku":
        _add_generic(
            document, entry, _puzzle_title(entry), new_page, solved=False, bookmark=bookmark
        )
        return
    text = entry.text
    presentation = load_presentation(text)
    _title(document, _puzzle_title(entry), new_page, bookmark, entry.section)
    document.add_heading(phrase(text, "doc.rules_title"), level=2)
    for line in presentation.intro or [phrase(text, "doc.rules")]:
        _runs(document.add_paragraph(style="List Bullet"), line)
    notes = board_notes(text)
    for note in notes:
        document.add_paragraph(note, style="List Bullet")
    note = phrase(text, "doc.murdoku_note")
    _runs(document.add_paragraph(style="List Bullet"), note)
    clues = clue_paragraphs(text, presentation.clue_text)
    _, aspect = _picture_stream(text, None, None)
    clue_lines = sum(-(-len(clue) // CLUE_LINE_CHARS) for clue in clues)
    rules = [*(presentation.intro or [phrase(text, "doc.rules")]), *notes, note]
    rules_lines = sum(-(-len(line) // CLUE_LINE_CHARS) for line in rules)
    rules_lines += _list_lines(text)
    width = _puzzle_board_width(aspect, clue_lines, len(clues), rules_lines)
    backdrop = width > MIN_BOARD_WIDTH_CM
    if not backdrop:
        _, aspect = _picture_stream(text, None, None, backdrop=False)
        width = _puzzle_board_width(aspect, clue_lines, len(clues), rules_lines)
    _board_and_figures(document, entry, presentation, width, backdrop)
    _board_lists(document, text)
    document.add_heading(phrase(text, "doc.clues"), level=2)
    for number, clue in enumerate(clues, 1):
        _numbered(document, number, clue, clue_lines > MANY_CLUES)
    if presentation.question:
        _boxed(document, presentation.question)


def _add_solution(document: Any, entry: Entry, title: str, bookmark: str | None = None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    if entry.kind != "murdoku":
        _add_generic(document, entry, title, True, solved=True, bookmark=bookmark)
        return
    text = entry.text
    presentation = load_presentation(text)
    _title(document, title, True, bookmark, entry.section)
    stream, aspect = _picture_stream(text, entry.solution, entry.role_statuses)
    blocks = [b.strip().removeprefix("- ") for b in explain(text, numbered=False).split("\n\n")]
    blocks = [b for b in blocks if b]
    lines = sum(-(-len(b) // CLUE_LINE_CHARS) for b in blocks)
    small = lines > MANY_SOLUTION_LINES
    width = _puzzle_board_width(aspect, int(lines * (0.8 if small else 1)), 0, 6)
    if width <= max(MIN_BOARD_WIDTH_CM, SOLUTION_MIN_BOARD_CM):
        stream, aspect = _picture_stream(text, entry.solution, entry.role_statuses, False)
        width = _puzzle_board_width(aspect, int(lines * (0.8 if small else 1)), 0, 6)
    document.add_picture(stream, width=Cm(max(width, SOLUTION_MIN_BOARD_CM)))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading(phrase(text, "doc.how"), level=2)
    for block in blocks:
        paragraph = document.add_paragraph(style="List Bullet")
        _runs(paragraph, block)
        if small:
            for run in paragraph.runs:
                run.font.size = Pt(8)
    for line in presentation.reveal:
        _boxed(document, line)


AFTER_SPACING = (
    "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
    "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
)  # fmt: skip
AFTER_BORDERS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)  # fmt: skip


def _bookmark(paragraph: Any, name: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(abs(hash(name)) % 100000))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), start.get(qn("w:id")))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _contents(document: Any, heading: str, items: list[tuple[str, str, str]]) -> None:
    """A table of contents: each line links to its bookmark; Word fills in the page number.

    Items are (section, label, bookmark); a new section gets a small heading.
    """
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document.add_heading(heading, level=1)
    current = ""
    for section, label, name in items:
        if section and section != current:
            current = section
            title = document.add_paragraph()
            title.paragraph_format.space_before = Pt(8)
            title.paragraph_format.keep_with_next = True
            _section_run(title.add_run(section))
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), name)
        link.set(qn("w:history"), "1")
        paragraph._p.append(link)

        def run(child: Any, link: Any = link) -> None:
            holder = OxmlElement("w:r")
            holder.append(child)
            link.append(holder)

        text = OxmlElement("w:t")
        text.text = label
        run(text)
        run(OxmlElement("w:tab"))
        for kind in ("begin", None, "separate", "result", "end"):
            if kind is None:
                code = OxmlElement("w:instrText")
                code.set(qn("xml:space"), "preserve")
                code.text = f" PAGEREF {name} \\h "
                run(code)
            elif kind == "result":
                shown = OxmlElement("w:t")
                shown.text = "·"
                run(shown)
            else:
                mark = OxmlElement("w:fldChar")
                mark.set(qn("w:fldCharType"), kind)
                if kind == "begin":
                    mark.set(qn("w:dirty"), "true")
                run(mark)


def _section_run(run: Any) -> None:
    """Format a run like a section heading: bold, spaced capitals."""
    from docx.shared import Pt, RGBColor

    run.bold = True
    run.font.all_caps = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    _letter_spacing(run._element.get_or_add_rPr(), 30)


def _letter_spacing(properties: Any, twentieths: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twentieths))
    properties.insert_element_before(spacing, *AFTER_SPACING)


def _add_borders(properties: Any, borders: Any) -> None:
    properties.insert_element_before(borders, *AFTER_BORDERS)


def _plain_font(style: Any, size: int, bold: bool, color: str) -> None:
    """Calibri in `color`, dropping the template's theme fonts and colour."""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = "Calibri"
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attribute}"), None)
    fonts.set(qn("w:eastAsia"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def _style_headings(document: Any) -> None:
    """Title and section styles of the original hand-made puzzle documents."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    title = document.styles["Heading 1"]
    _plain_font(title, 20, True, "222222")
    title.font.all_caps = True
    _letter_spacing(title.element.get_or_add_rPr(), 100)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = False
    section = document.styles["Heading 2"]
    _plain_font(section, 11, True, "333333")
    section.font.all_caps = True
    _letter_spacing(section.element.get_or_add_rPr(), 30)
    section.paragraph_format.space_before = Pt(9)
    section.paragraph_format.space_after = Pt(4)
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", "BBBBBB")):
        edge.set(qn(f"w:{key}"), value)
    borders.append(edge)
    _add_borders(section.element.get_or_add_pPr(), borders)


COPYRIGHT_NOTICE = (
    "\u00a9 Otto Perdeck. All rights reserved. No copying, redistribution or other use "
    "of these puzzles without the explicit written permission of the author."
)


def _footer(document: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.text = COPYRIGHT_NOTICE
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(7)


def _compact(document: Any) -> None:
    """Small margins and text so a puzzle, and its solution, each fit on one page."""
    from docx.shared import Cm, Pt

    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(1.4)
        section.left_margin = section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    _style_headings(document)
    normal.paragraph_format.space_after = Pt(3)
    for name in ("List Number", "List Bullet"):
        document.styles[name].paragraph_format.space_after = Pt(1)


def write_booklet(entries: list[Entry], path: Path) -> None:
    """Write one document: every puzzle on its own page, then the solved ones' solutions.

    A figure image path is relative to its puzzle file; a missing image is skipped,
    leaving the text-only entry. Raises ImportError without python-docx.
    """
    from docx import Document

    document = Document()
    _compact(document)
    _footer(document)
    solved = [entry for entry in entries if entry.solution is not None]
    if len(entries) > 1:
        items = [
            (entry.section, _puzzle_title(entry), f"puzzle{n}") for n, entry in enumerate(entries)
        ]
        if solved:
            items.append(("", phrase(solved[0].text, "doc.solutions"), "solution0"))
        _contents(document, phrase(entries[0].text, "doc.contents"), items)
    for number, entry in enumerate(entries):
        _add_puzzle(
            document, entry, new_page=number > 0 or len(entries) > 1, bookmark=f"puzzle{number}"
        )
    for number, entry in enumerate(solved):
        _add_solution(document, entry, _solution_title(entry), f"solution{number}")
    document.save(str(path))
